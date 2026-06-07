import argparse
import base64
import json
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

from aigc2d_ocr import parse_combined_text, replace_combined_page


SYSTEM_PROMPT = """你是专业的日语漫画汉化翻译。
将输入的日语漫画 OCR 文本翻译为自然、简洁的简体中文。
严格使用提供的术语表译名，不得自行改写已有对应关系。
保留原文的对话顺序、换行、拟声词、标点和语气；不要添加解释、注释、Markdown 或页码。
如果单独一行假名只是相邻汉字行的普通读音，请完全忽略该假名行，只翻译包含汉字的正文行。
纯假名的正文、助词、语气词、拟声词仍须正常翻译，不得因为没有汉字而忽略。
不要逐行分析、解释翻译过程、复述要求或重复原文；将跨行组成的一句话合并为自然中文。
输出前删除由注音造成的重复译文。例如相邻的“感想 / かんそう”只能输出一次“感想”。
只在 JSON 的 translation 字段中输出最终中文译文。"""

TRANSLATION_FORMAT = {
    "type": "object",
    "properties": {"translation": {"type": "string"}},
    "required": ["translation"],
}

VL_REVIEW_FORMAT = {
    "type": "object",
    "properties": {
        "segments": {
            "type": "array",
            "items": {"type": "string"},
        }
    },
    "required": ["segments"],
}

VL_REVIEW_PROMPT = """你是专业的漫画汉化校对。
请结合漫画图片、日语 OCR 原文和中文初译，输出校对后的简体中文译文。
修正错译、漏译、人物语气和标点；严格沿用提供的术语表。
初译与原文意思一致时不要擅自改变含义，只调整断句和表达。
不要沿用 OCR 为适应竖排和气泡宽度产生的断行。每个对话框中的完整句子必须合并为一行。
不同对话框、不同说话者、标题、拟声词之间使用换行或空行分隔，不要把整页合并成一段。
输出前逐段检查：如果某段中的连续行能够组成一句语法完整的话，必须删除这些行间换行。
忽略普通注音，不要重复译文，不要添加解释、Markdown、页码或原文。
将标题、每个对话框和每个独立拟声词分别放入 JSON 的 segments 数组；每个数组元素内部不得换行。"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Translate page-marked Japanese OCR text with local Qwen3-30B-A3B."
    )
    parser.add_argument("input", type=Path, help="Combined OCR text containing P{page} markers")
    parser.add_argument("--pages", help="Inclusive page range, for example 7-56")
    parser.add_argument(
        "--review-page",
        type=int,
        help="Redo one page and replace it in the existing --combined-output file",
    )
    parser.add_argument(
        "--glossary", type=Path, default=Path("名词表.xlsx"), help="Excel glossary path"
    )
    parser.add_argument("--sheet", help="Glossary worksheet name (default: first sheet)")
    parser.add_argument(
        "--output-dir", type=Path, default=Path("translation_text"), help="Per-page output"
    )
    parser.add_argument(
        "--combined-output",
        type=Path,
        default=Path("translation_text/pages_translated.txt"),
        help="Combined translation path",
    )
    parser.add_argument("--model", default="qwen3:30b", help="Ollama model name")
    parser.add_argument("--base-url", default="http://localhost:11434", help="Ollama URL")
    parser.add_argument(
        "--image-dir",
        type=Path,
        default=Path("exported_jpg"),
        help="Directory containing page_NNNN.jpg for VL proofreading",
    )
    parser.add_argument(
        "--vl-model",
        default="qwen3-vl:30b-a3b-instruct",
        help="Ollama vision model used for proofreading",
    )
    parser.add_argument(
        "--vl-base-url",
        default="http://localhost:11436",
        help="Ollama URL for VL proofreading",
    )
    parser.add_argument(
        "--no-vl-review",
        action="store_true",
        help="Skip the Qwen3-VL proofreading stage",
    )
    parser.add_argument(
        "--vl-review-only",
        action="store_true",
        help="Proofread existing --combined-output translations without retranslating",
    )
    parser.add_argument("--temperature", type=float, default=0.1)
    parser.add_argument("--num-predict", type=int, default=4096)
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument(
        "--request-timeout",
        type=int,
        default=300,
        help="Timeout in seconds for each translation or VL proofreading request",
    )
    parser.add_argument("--resume", action="store_true", help="Reuse existing page translations")
    parser.add_argument(
        "--keep-page-files",
        action="store_true",
        help="Keep per-page translation files after the combined output is complete",
    )
    parser.add_argument("--dry-run", action="store_true", help="Show pages and glossary only")
    return parser.parse_args()


def filter_pages(pages: list[tuple[int, str]], page_range: str | None) -> list[tuple[int, str]]:
    if not page_range:
        return pages
    try:
        start_text, end_text = page_range.split("-", maxsplit=1)
        start, end = int(start_text), int(end_text)
    except ValueError as exc:
        raise ValueError("--pages must use the format START-END, for example 7-56") from exc
    if start < 1 or end < start:
        raise ValueError("--pages must be a positive range with START <= END")
    selected = [(page, text) for page, text in pages if start <= page <= end]
    if not selected:
        raise ValueError(f"No input pages matched --pages {page_range}")
    return selected


def load_glossary(path: Path, sheet_name: str | None) -> tuple[list[tuple[str, str]], list[str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Glossary does not exist: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))

    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if "日语名" in row and "中文名" in row
        ),
        None,
    )
    if header_index is None:
        raise ValueError("Glossary must contain headers: 日语名 and 中文名")

    headers = list(rows[header_index])
    japanese_index = headers.index("日语名")
    chinese_index = headers.index("中文名")
    skills_index = headers.index("角色的技能名") if "角色的技能名" in headers else None

    terms: list[tuple[str, str]] = []
    skills: list[str] = []
    for row in rows[header_index + 1 :]:
        japanese = row[japanese_index] if japanese_index < len(row) else None
        chinese = row[chinese_index] if chinese_index < len(row) else None
        if japanese and chinese:
            terms.append((str(japanese).strip(), str(chinese).strip()))
        if skills_index is not None and skills_index < len(row) and row[skills_index]:
            skills.extend(
                line.strip()
                for line in str(row[skills_index]).splitlines()
                if line.strip()
            )
    return terms, skills


def glossary_prompt(source: str, terms: list[tuple[str, str]], skills: list[str]) -> str:
    relevant = [(jp, zh) for jp, zh in terms if jp in source]
    lines = [f"{jp} => {zh}" for jp, zh in relevant]
    if skills:
        lines.append("既有技能译名参考：" + "；".join(skills))
    return "\n".join(lines) if lines else "本页没有匹配到术语表条目。"


def strip_thinking(text: str) -> str:
    stripped = text.strip()
    if "</think>" in stripped.lower():
        closing_index = stripped.lower().rfind("</think>")
        return stripped[closing_index + len("</think>") :].strip()
    return re.sub(r"(?is)^\s*<think>.*?</think>\s*", "", stripped).strip()


def request_translation(
    source: str,
    terms: list[tuple[str, str]],
    skills: list[str],
    model: str,
    base_url: str,
    temperature: float,
    num_predict: int,
    request_timeout: int,
) -> str:
    prompt = (
        f"/no_think\n术语表：\n{glossary_prompt(source, terms, skills)}\n\n"
        f"待翻译日文：\n{source}"
    )
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        "think": False,
        "stream": False,
        "format": TRANSLATION_FORMAT,
        "options": {"temperature": temperature, "num_predict": num_predict},
    }
    request = urllib.request.Request(
        f"{base_url.rstrip('/')}/api/chat",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=request_timeout) as response:
            result = json.load(response)
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Ollama returned HTTP {exc.code}: {details}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Could not connect to Ollama: {exc.reason}") from exc
    except TimeoutError as exc:
        raise RuntimeError(
            f"Translation request timed out after {request_timeout} seconds"
        ) from exc

    try:
        message = result["message"]
        content = strip_thinking(message.get("content") or "")
        translated = json.loads(content)["translation"]
        if not isinstance(translated, str):
            raise TypeError("translation is not a string")
        return translated.strip()
    except (json.JSONDecodeError, KeyError, TypeError) as exc:
        raise RuntimeError(f"Unexpected Ollama response: {json.dumps(result)}") from exc


def request_vl_review(
    image_path: Path,
    source: str,
    draft: str,
    terms: list[tuple[str, str]],
    skills: list[str],
    args: argparse.Namespace,
) -> str:
    prompt = (
        f"/no_think\n术语表：\n{glossary_prompt(source, terms, skills)}\n\n"
        f"日语 OCR 原文：\n{source}\n\n中文初译：\n{draft}"
    )
    payload = {
        "model": args.vl_model,
        "messages": [
            {"role": "system", "content": VL_REVIEW_PROMPT},
            {
                "role": "user",
                "content": prompt,
                "images": [base64.b64encode(image_path.read_bytes()).decode("ascii")],
            },
        ],
        "think": False,
        "stream": False,
        "format": VL_REVIEW_FORMAT,
        "options": {"temperature": args.temperature, "num_predict": args.num_predict},
    }
    request = urllib.request.Request(
        f"{args.vl_base_url.rstrip('/')}/api/chat",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=args.request_timeout) as response:
            result = json.load(response)
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"VL proofreading returned HTTP {exc.code}: {details}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Could not connect to VL proofreading Ollama: {exc.reason}") from exc
    except TimeoutError as exc:
        raise RuntimeError(
            f"VL proofreading request timed out after {args.request_timeout} seconds"
        ) from exc

    try:
        content = strip_thinking(result["message"].get("content") or "")
        segments = json.loads(content)["segments"]
        if not isinstance(segments, list) or not all(
            isinstance(segment, str) for segment in segments
        ):
            raise TypeError("segments is not a string array")
        cleaned = [
            re.sub(r"[\r\n]+", "", segment).strip()
            for segment in segments
            if segment.strip()
        ]
        return "\n\n".join(cleaned)
    except (json.JSONDecodeError, KeyError, TypeError) as exc:
        raise RuntimeError(f"Unexpected VL proofreading response: {json.dumps(result)}") from exc


def invalid_translation_reason(text: str, source: str | None = None) -> str | None:
    stripped = text.strip()
    if not stripped:
        return "empty translation"
    if "<think>" in stripped.lower():
        return "thinking output detected"
    if re.search(r"(?m)^P\d+\s*$", stripped):
        return "unexpected page marker detected"
    analysis_markers = (
        "首先，用户要求",
        "我需要翻译",
        "分析文本",
        "翻译步骤",
        "最终翻译",
        "现在，写输出",
        "提供的术语表",
    )
    if any(marker in stripped for marker in analysis_markers):
        return "analysis or explanation detected"
    if source and len(stripped) > max(1000, len(source) * 8):
        return (
            f"translation is too long relative to source "
            f"({len(stripped)} vs {len(source)} characters)"
        )
    if len(stripped) > 10000:
        return f"translation is abnormally long ({len(stripped)} characters)"
    return None


def translate_with_retries(
    source: str,
    terms: list[tuple[str, str]],
    skills: list[str],
    args: argparse.Namespace,
) -> str:
    last_reason = "unknown invalid translation"
    for attempt in range(1, args.retries + 1):
        text = request_translation(
            source,
            terms,
            skills,
            args.model,
            args.base_url,
            args.temperature,
            args.num_predict,
            args.request_timeout,
        )
        reason = invalid_translation_reason(text, source)
        if not reason:
            return text.strip()
        last_reason = reason
        print(f"Invalid translation ({attempt}/{args.retries}): {reason}", file=sys.stderr)
    raise RuntimeError(f"Translation failed after {args.retries} attempts: {last_reason}")


def review_translation(
    page: int,
    source: str,
    draft: str,
    terms: list[tuple[str, str]],
    skills: list[str],
    args: argparse.Namespace,
) -> str:
    image_path = args.image_dir / f"page_{page:04d}.jpg"
    if not image_path.is_file():
        raise FileNotFoundError(f"VL proofreading image does not exist: {image_path}")
    print(f"VL proofreading P{page}")
    text = request_vl_review(image_path, source, draft, terms, skills, args)
    reason = invalid_translation_reason(text, source)
    if reason:
        raise RuntimeError(f"VL proofreading produced invalid translation: {reason}")
    return text


def write_combined(path: Path, translations: list[tuple[int, str]]) -> None:
    content = "\n\n".join(f"P{page}\n{text.strip()}" for page, text in translations) + "\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def main() -> int:
    args = parse_args()
    try:
        input_pages = parse_combined_text(args.input.read_text(encoding="utf-8"))
        if args.review_page is not None:
            if args.review_page < 1:
                raise ValueError("--review-page must be greater than zero")
            if args.pages:
                raise ValueError("--review-page cannot be combined with --pages")
            if not args.combined_output.is_file():
                raise FileNotFoundError(
                    f"--review-page requires an existing combined output: {args.combined_output}"
                )
            pages = [(page, text) for page, text in input_pages if page == args.review_page]
            if not pages:
                raise ValueError(f"Input does not contain P{args.review_page}: {args.input}")
        else:
            pages = filter_pages(input_pages, args.pages)
        terms, skills = load_glossary(args.glossary, args.sheet)
        print(
            f"Selected {len(pages)} page(s): P{pages[0][0]} through P{pages[-1][0]}; "
            f"loaded {len(terms)} term mapping(s)"
        )
        if args.dry_run:
            for page, source in pages:
                matched = sum(1 for japanese, _ in terms if japanese in source)
                print(f"P{page}: {len(source)} source characters, {matched} matched terms")
            return 0
        if args.num_predict < 1 or args.retries < 1 or args.request_timeout < 1:
            raise ValueError(
                "--num-predict, --retries and --request-timeout must be greater than zero"
            )
        existing_translations = {}
        existing_translation_pages = []
        if args.vl_review_only:
            if args.no_vl_review:
                raise ValueError("--vl-review-only cannot be combined with --no-vl-review")
            if args.review_page is not None:
                raise ValueError("--vl-review-only cannot be combined with --review-page")
            if not args.combined_output.is_file():
                raise FileNotFoundError(
                    f"--vl-review-only requires an existing combined output: "
                    f"{args.combined_output}"
                )
            existing_translation_pages = parse_combined_text(
                args.combined_output.read_text(encoding="utf-8")
            )
            existing_translations = dict(existing_translation_pages)
            missing = [page for page, _ in pages if page not in existing_translations]
            if missing:
                raise ValueError(
                    f"Combined translation is missing page(s): "
                    f"{', '.join(f'P{page}' for page in missing)}"
                )

        args.output_dir.mkdir(parents=True, exist_ok=True)
        translations = []
        for page, source in pages:
            output_path = args.output_dir / f"page_{page:04d}.txt"
            if args.vl_review_only:
                print(f"VL proofreading existing P{page}")
                text = review_translation(
                    page, source, existing_translations[page], terms, skills, args
                )
                if args.keep_page_files:
                    output_path.write_text(text + "\n", encoding="utf-8")
                    print(f"Saved: {output_path}")
            elif args.resume and args.review_page is None and output_path.is_file():
                text = output_path.read_text(encoding="utf-8").strip()
                reason = invalid_translation_reason(text, source)
                if reason:
                    print(f"Redoing invalid translation: {output_path} ({reason})")
                    text = translate_with_retries(source, terms, skills, args)
                    if not args.no_vl_review:
                        text = review_translation(page, source, text, terms, skills, args)
                    output_path.write_text(text + "\n", encoding="utf-8")
                else:
                    print(f"Reused: {output_path}")
            else:
                print(f"Translating P{page}")
                text = translate_with_retries(source, terms, skills, args)
                if not args.no_vl_review:
                    text = review_translation(page, source, text, terms, skills, args)
                output_path.write_text(text + "\n", encoding="utf-8")
                print(f"Saved: {output_path}")

            translations.append((page, text))
            if args.vl_review_only:
                existing_translations[page] = text
                write_combined(
                    args.combined_output,
                    [
                        (number, existing_translations[number])
                        for number, _ in existing_translation_pages
                    ],
                )
                if not args.keep_page_files:
                    output_path.unlink(missing_ok=True)
                print(f"Updated P{page} in combined translation")
            elif args.review_page is None:
                write_combined(args.combined_output, translations)

        if args.review_page is not None:
            replace_combined_page(args.combined_output, args.review_page, translations[0][1])
            print(
                f"Replaced P{args.review_page} in combined translation: "
                f"{args.combined_output}"
            )
        else:
            print(f"Saved combined translation: {args.combined_output}")
        if not args.keep_page_files:
            for page, _ in pages:
                (args.output_dir / f"page_{page:04d}.txt").unlink(missing_ok=True)
            print(f"Deleted {len(pages)} per-page intermediate translation file(s)")
    except (OSError, RuntimeError, ValueError, KeyError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
